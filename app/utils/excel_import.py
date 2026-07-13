"""Importador flexible de listas de precios (Excel .xlsx/.xlsm o CSV).

En vez de asumir columnas fijas ('ARTICULOS'/'PRECIO'), detecta la fila de
encabezados de la planilla y permite mapear qué columna corresponde a cada
campo (nombre, precio, código). El talle se contempla en el mapeo pero recién
se guarda como dato propio en la Fase 3; por ahora, si se mapea, se agrega al
nombre del producto.

Uso desde la UI:
    hojas = listar_hojas(path)
    headers, filas = leer_hoja(path, hoja)
    mapeo = sugerir_mapeo(headers)          # {campo: indice_columna | None}
    items = filas_a_items(filas, mapeo)     # [{'nombre','codigo','pvp'}]
    db.upsert_conceptos(items)
"""

import csv

import openpyxl

from utils.helpers import separar_codigo

# Campos que se pueden mapear. 'nombre' y 'precio' son obligatorios.
CAMPOS = ("nombre", "precio", "codigo", "talle")

# Pistas para auto-detectar cada campo por el texto del encabezado.
_PISTAS = {
    "nombre": ("articulo", "artículo", "producto", "descrip", "detalle", "nombre", "item"),
    "precio": ("precio", "pvp", "importe", "valor", "monto"),
    "codigo": ("codigo", "código", "cod.", "cod ", "sku"),
    "talle": ("talle", "talla", "medida", "size"),
}

_CSV_HOJA = "CSV"


def _norm(v) -> str:
    return str(v if v is not None else "").strip().lower()


def _es_fila_encabezado(celdas) -> bool:
    """Una fila es encabezado si tiene al menos 2 celdas de texto no numérico."""
    textos = 0
    for c in celdas:
        if isinstance(c, str) and c.strip() and not c.strip().replace(",", "").replace(".", "").isdigit():
            textos += 1
    return textos >= 2


def listar_hojas(path: str) -> list[str]:
    if path.lower().endswith(".csv"):
        return [_CSV_HOJA]
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _filas_crudas(path: str, hoja: str | None) -> list[list]:
    if path.lower().endswith(".csv"):
        # Detecta separador (coma o punto y coma) con el Sniffer.
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
            muestra = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(muestra, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return [list(r) for r in csv.reader(fh, dialect)]
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[hoja] if hoja and hoja in wb.sheetnames else wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def leer_hoja(path: str, hoja: str | None = None) -> tuple[list[str], list[list]]:
    """Devuelve (encabezados, filas_de_datos) de la hoja. Busca la fila de
    encabezados entre las primeras filas (saltando membretes) y devuelve las
    filas siguientes como datos."""
    filas = _filas_crudas(path, hoja)
    header_idx = None
    for i, fila in enumerate(filas[:40]):
        if _es_fila_encabezado(fila):
            header_idx = i
            break
    if header_idx is None:
        # Sin encabezado claro: usar columnas genéricas.
        ancho = max((len(f) for f in filas), default=0)
        headers = [f"Columna {c + 1}" for c in range(ancho)]
        return headers, filas
    headers = [str(c).strip() if c is not None else "" for c in filas[header_idx]]
    headers = [h or f"Columna {i + 1}" for i, h in enumerate(headers)]
    return headers, filas[header_idx + 1:]


def sugerir_mapeo(headers: list[str]) -> dict:
    """Auto-detecta qué columna es cada campo por el nombre del encabezado.
    Devuelve {campo: indice | None}."""
    mapeo = {campo: None for campo in CAMPOS}
    for idx, h in enumerate(headers):
        hn = _norm(h)
        if not hn:
            continue
        for campo, pistas in _PISTAS.items():
            if mapeo[campo] is None and any(p in hn for p in pistas):
                mapeo[campo] = idx
    return mapeo


def _celda(fila: list, idx) -> str:
    if idx is None or idx < 0 or idx >= len(fila):
        return ""
    v = fila[idx]
    return str(v).strip() if v is not None else ""


def _to_float(texto: str) -> float | None:
    s = texto.strip()
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    s = "".join(c for c in s if c.isdigit() or c == ".")
    try:
        return float(s)
    except ValueError:
        return None


def filas_a_items(filas: list[list], mapeo: dict) -> list[dict]:
    """Aplica el mapeo a las filas y devuelve items {'nombre','codigo','pvp'}
    listos para db.upsert_conceptos. Descarta filas sin nombre o sin precio
    válido."""
    out = []
    for fila in filas:
        nombre = _celda(fila, mapeo.get("nombre"))
        if not nombre or nombre.startswith("*") or "CONSULTAR" in nombre.upper():
            continue
        pvp = _to_float(_celda(fila, mapeo.get("precio")))
        if pvp is None:
            continue
        codigo = _celda(fila, mapeo.get("codigo"))
        talle = _celda(fila, mapeo.get("talle"))
        if not codigo:
            # Si no hay columna de código, intentar extraerlo del nombre.
            nombre, codigo = separar_codigo(nombre)
        if talle:
            nombre = f"{nombre} {talle}".strip()
        out.append({"nombre": nombre, "codigo": codigo, "pvp": pvp})
    return out
